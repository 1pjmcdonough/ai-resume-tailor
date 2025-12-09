import streamlit as st
import io
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from homepage import render_header
import copy


def main():
    render_header()

    if "cover_letter" in st.session_state:
        if "cover_letter_edits" not in st.session_state:
            st.session_state.cover_letter_edits = copy.deepcopy(st.session_state.cover_letter)

        tab1, tab2 = st.tabs(["Preview", "Edit"])
        
        with tab1:
            display_preview()
            download_cl(st.session_state.cover_letter_edits)
        
        with tab2:
            create_editable_form()
    else:
        st.info("Please generate a cover letter first.")


def display_preview():
    """Displays a preview of the cover letter with intro, T-table pairs, and closing paragraph."""
    st.subheader("Preview Cover Letter")

    st.write(st.session_state.cover_letter_edits["intro_paragraph"])
    
    # T-table header
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("##### Your Requirements")
    with col2:
        st.markdown("##### My Qualifications")
    st.divider()
    
    # Display each requirement-qualification pair
    for pair in st.session_state.cover_letter_edits["t_table"]:
        col1, col2 = st.columns(2)
        with col1:
            st.write(pair["job_requirement"])
        with col2:
            st.write(pair["my_qualification"])
        st.divider()

    st.write(st.session_state.cover_letter_edits["closing_paragraph"])


def download_cl(edits: dict):
    """Creates a download button for the cover letter as a Word document.
    
    Args:
        edits (dict): The cover letter data dictionary containing intro, T-table, and closing paragraph.
    """
    doc_bytes = create_ttable_cover_letter_bytes(edits)
        
    st.download_button(
        label="⬇ Download Cover Letter",
        data=doc_bytes,
        file_name="cover_letter.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary"
    )


def create_ttable_cover_letter_bytes(edits: dict):
    """Creates a Word document from the cover letter edits and returns it as bytes.
    
    Args:
        edits (dict): The cover letter data dictionary containing intro, T-table, and closing paragraph.
    
    Returns:
        bytes: The Word document as a byte string ready for download.
    """
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    
    intro_para = doc.add_paragraph()
    intro_para.add_run(edits["intro_paragraph"])
    
    doc.add_paragraph()
    
    if edits["t_table"]:
        table = doc.add_table(rows=len(edits["t_table"])+1, cols=2) # +1 for the header
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Set column widths
        table.columns[0].width = Inches(3.5)
        table.columns[1].width = Inches(3.5)
        
        header_cells = table.rows[0].cells
        header_cells[0].text = "Your Requirements"
        header_cells[1].text = "My Qualifications"
        
        for cell in header_cells:
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                
        for i, item in enumerate(edits["t_table"], 1):
            row_cells = table.rows[i].cells
            row_cells[0].text = item["job_requirement"]
            row_cells[1].text = item["my_qualification"]
            
            for cell in row_cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph()  # Add some space after the table
    
    closing_para = doc.add_paragraph()
    closing_para.add_run(edits["closing_paragraph"])
    
    # Convert document to bytes
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    doc_buffer.seek(0)
    
    return doc_buffer.getvalue()


def create_editable_form():
    """Creates an editable form for the cover letter data with text areas and row management."""
    st.subheader("Edit Cover Letter")

    # Introduction paragraph editor
    st.session_state.cover_letter_edits["intro_paragraph"] = st.text_area(
        label="Introduction Paragraph", 
        value=st.session_state.cover_letter_edits["intro_paragraph"],
        height=180,
        key="intro_edit"
    )

    # T-table header
    col1, col2, _ = st.columns([2, 2, 0.3])
    with col1:
        st.markdown("##### Your Requirements")
    with col2:
        st.markdown("##### My Qualifications")
     
    # Display each requirement-qualification pair with edit and delete options
    for i, item in enumerate(st.session_state.cover_letter_edits["t_table"]):
        col1, col2, col3 = st.columns([2, 2, 0.3])

        with col1:
            item['job_requirement'] = st.text_area(
                "Job Requirement",
                value=item["job_requirement"],
                height=80,
                key=f"req_{i}",
                label_visibility="collapsed"
            )

        with col2:
            item['my_qualification'] = st.text_area(
                "My Qualification",
                value=item['my_qualification'],
                height=80,
                key=f"qual_{i}",
                label_visibility="collapsed"
            )

        with col3:
            st.write("")
            if st.button("🗑️", key=f"del_{i}", help="Delete this row"):
                st.session_state.cover_letter_edits["t_table"] = [
                    row for j, row in enumerate(st.session_state.cover_letter_edits["t_table"]) if j != i
                ]
                st.rerun()

    # Add new row button
    if st.button("➕ Add Row", type="secondary"):
        new_row = {'job_requirement': '', 'my_qualification': ''}
        st.session_state.cover_letter_edits['t_table'].append(new_row)
        st.rerun()

    st.markdown("---")
    
    # Closing paragraph editor
    st.session_state.cover_letter_edits['closing_paragraph'] = st.text_area(
        label="Closing Paragraph", 
        value=st.session_state.cover_letter_edits['closing_paragraph'],
        height=180,
        key="closing_edit"
    )

    # Action buttons
    st.markdown("### Actions")
    col1, col2 = st.columns(2)
    with col1:
        if st.button("💾 Save Changes", type="primary"):
            st.rerun()

    with col2:
        if st.button("↩ Reset Changes"):
            st.session_state.cover_letter_edits = copy.deepcopy(st.session_state.cover_letter)
            st.info("Edits reset to original cover letter.")
            st.rerun()


if __name__ == "__main__":
    main()