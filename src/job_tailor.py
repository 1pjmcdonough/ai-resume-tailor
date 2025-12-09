"""
JobTailor Module

This module provides a JobTailor class that uses AI to tailor resumes and generate
cover letters based on job descriptions. It supports parsing PDF, DOCX, and TXT files,
and saves generated responses with metadata for tracking.
"""

import json
from pathlib import Path
from typing import Optional

import PyPDF2
from docx import Document
from openai import OpenAI

from src.helpers import Prompt, set_client


class JobTailor:
    """
    A class to tailor resumes and generate cover letters using AI based on job descriptions.
    """
    
    def __init__(self):
        """
        Initialize the JobTailor instance with default settings.
        """
        self.model: Optional[str] = "grok-4-0709"  # AI model to use for generation
        self.client: OpenAI = set_client(self.model)  # OpenAI client instance
        self.user_context: Optional[str] = ""  # Additional user-provided context
        self.resume_filename: Optional[str] = ""  # Name of the uploaded resume file
        self.job_desc_filename: Optional[str] = ""  # Name of the uploaded job description file


    def upload_resume(self, resume_path: Path) -> None:
        """
        Upload and parse a resume file.
        
        Args:
            resume_path: Path to the resume file (PDF, DOCX, or TXT)
        
        The parsed text is stored in self.resume_text and the filename in self.resume_filename.
        """
        self.resume_text = self.parse_file(resume_path)
        self.resume_filename = resume_path.name

    
    def upload_job_desc(self, job_desc_path: Path) -> None:
        """
        Upload and parse a job description file.
        
        Args:
            job_desc_path: Path to the job description file (PDF, DOCX, or TXT)
        
        The parsed text is stored in self.job_desc_text and the filename in self.job_desc_filename.
        """
        self.job_desc_text = self.parse_file(job_desc_path)
        self.job_desc_filename = job_desc_path.name


    def parse_file(self, file_path: Path) -> str:
        """
        Extract text content from a file. 
        
        Args:
            file_path: Path to the file to parse
        
        Returns:
            str: The extracted text content from the file
        
        Raises:
            ValueError: If the file type is not supported (not .pdf, .docx, or .txt)
        """
        file_type = file_path.suffix.lower()

        if file_type == ".pdf":
            # Extract text from all pages in the PDF
            pdf_reader = PyPDF2.PdfReader(file_path)
            return "\n".join([p.extract_text() for p in pdf_reader.pages])
        elif file_type == ".docx":
            # Extract text from all paragraphs in the Word document
            doc = Document(file_path)
            return "\n".join([p.text for p in doc.paragraphs])
        elif file_type == ".txt":
            # Read plain text file directly
            with open(file_path, "r") as f:
                 return f.read()
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    
    def set_model(self, model: str) -> None:
        """
        Set the AI model to use for generation.
        
        Args:
            model: Name of the model to use
        
        Note: After setting a new model, call configure_client() to update the client.
        """
        self.model = model


    def configure_client(self) -> None:
        """
        Reconfigure the OpenAI client with the current model setting.
        
        This should be called after changing the model using set_model() to ensure
        the client is using the correct model configuration.
        """
        self.client = set_client(self.model)

        
    def tailor_resume(self) -> dict:
        """
        Generate tailored resume edits based on the uploaded resume and job description.
        
        Uses AI to analyze the resume and job description, then generates suggestions
        for how to tailor the resume to better match the job requirements.
        
        Returns:
            dict: JSON response containing the tailored resume edits/suggestions
        
        Raises:
            Exception: If there's an error during the tailoring process
        """
        try:
            resume_edits = self.generate_prompt("tailor_resume")
            self.save_response(resume_edits, "resume_edits")
            return resume_edits
        except Exception as e:
            raise Exception(f"Error tailoring resume: {str(e)}")
        
    
    def generate_cover_letter(self) -> dict:
        """
        Generate a cover letter based on the uploaded resume and job description.
        
        Uses AI to create a personalized cover letter that highlights relevant
        experience and skills from the resume that match the job requirements.
        
        Returns:
            dict: JSON response containing the generated cover letter
        
        Raises:
            Exception: If there's an error during the cover letter generation
        """
        try:
            cover_letter = self.generate_prompt("cover_letter")
            self.save_response(cover_letter, "cover_letter")
            return cover_letter
        except Exception as e:
            raise Exception(f"Error generating cover letter: {str(e)}")

        
    def generate_prompt(self, purpose: str) -> dict:
        """
        Generate AI response using prompts loaded from files.
        
        Loads system and user prompts for the specified purpose (e.g., "tailor_resume"
        or "cover_letter"), sends them to the AI model, and returns the parsed JSON response.
        
        Args:
            purpose: The purpose of the prompt (determines which prompt files to load)
                    Should match a folder name in the prompts directory
        
        Returns:
            dict: Parsed JSON response from the AI model
        
        The AI parameters are configured as:
        - temperature: 0.7 (moderate creativity)
        - top_p: 0.8 (nucleus sampling)
        - max_tokens: 4000 (maximum response length)
        """
        prompt = self.get_prompt(purpose)

        response = self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": prompt.sys_prompt},
                {"role": "user", "content": prompt.usr_prompt}
            ],
            temperature=0.7,
            top_p=0.8,
            max_tokens=4000
        )
        
        return json.loads(response.choices[0].message.content)
        

    def get_prompt(self, purpose: str) -> Prompt:
        """
        Load and prepare prompts for the specified purpose.
        
        Loads system and user prompts from the prompts directory, then replaces
        placeholders in the user prompt with actual resume text, job description,
        and user context.
        
        Args:
            purpose: The purpose/type of prompt to load (e.g., "tailor_resume", "cover_letter")
                    Must correspond to a folder name in the prompts directory
        
        Returns:
            Prompt: A Prompt object containing the system and user prompts
        
        The user prompt template supports the following placeholders:
        - {job_desc}: Replaced with the job description text
        - {resume}: Replaced with the resume text
        - {user_context}: Replaced with any additional user-provided context
        """
        sys_prompt_path = Path( Path(__file__).parent.parent, "prompts", purpose, "sys_prompt.txt" )
        usr_prompt_path = Path( Path(__file__).parent.parent, "prompts", purpose, "usr_prompt.txt" )

        with open(sys_prompt_path, "r") as f:
            sys_prompt = f.read()

        with open(usr_prompt_path, "r") as f:
            usr_prompt = f.read()
            usr_prompt = usr_prompt.replace("{job_desc}", self.job_desc_text)
            usr_prompt = usr_prompt.replace("{resume}", self.resume_text)
            usr_prompt = usr_prompt.replace("{user_context}", self.user_context)

        return Prompt(sys_prompt=sys_prompt, usr_prompt=usr_prompt)


    def save_response(self, response: dict, response_type: str) -> None:
        """
        Save a response and its metadata as JSON files in organized folders.
        
        Creates a directory structure in response_history/ organized by resume and job
        description filenames. Saves both the response content and metadata (filenames,
        user context, model used) for tracking and reference.
        
        Args:
            response: The response dictionary to save (e.g., tailored resume edits or cover letter)
            response_type: Type of response (e.g., "resume_edits", "cover_letter")
                         Used as the base filename for both the response and metadata files
        
        Raises:
            Exception: If there's an error during file saving
        
        Directory structure created:
        response_history/
          {resume_name}_for_{job_desc_name}/
            {response_type}.json
            {response_type}_metadata.json
        """
        try:
            response_history_path = Path(Path(__file__).parent.parent, "response_history")
            response_history_path.mkdir(exist_ok=True)
            
            curr_response_folder = f"{self.resume_filename.split('.')[0]}_for_{self.job_desc_filename.split('.')[0]}"
            curr_response_path = Path(response_history_path, curr_response_folder)
            curr_response_path.mkdir(exist_ok=True)

            filename = f"{response_type}.json"
            response_path = Path(curr_response_path, filename)
            with open(response_path, "w", encoding="utf-8") as f:
                json.dump(response, f, indent=2)
            
            metadata = {
                "resume_filename": self.resume_filename,
                "job_description_filename": self.job_desc_filename,
                "user_context": self.user_context,
                "model_used": self.model
            }
            
            metadata_path = Path(curr_response_path, f"{response_type}_metadata.json")
            with open(metadata_path, "w", encoding="utf-8") as f:
                json.dump(metadata, f, indent=2)
                
        except Exception as e:
            raise Exception(f"Error saving response: {str(e)}")